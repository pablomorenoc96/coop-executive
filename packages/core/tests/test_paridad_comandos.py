"""Comandos nuevos sin red ni modelo real: sitio web simulado, reunión, redacción, documentos y análisis."""
import json
from datetime import date, timedelta

import httpx
import pytest
import yaml
from click.testing import CliRunner
from docx import Document

from coopexecutive.cli import cli
from coopexecutive.crm import cases
from coopexecutive.documents import desde_markdown
from coopexecutive.documents.contenido import DocumentoPlano, Lista, Parrafo, Tabla, Titulo, a_markdown
from coopexecutive.grant_tools import matrix, reunion
from coopexecutive.grant_tools.matrix import PESOS
from coopexecutive.guardrails.sensibles import DATO_BANCARIO_OMITIDO
from coopexecutive.lectura import ErrorLectura
from coopexecutive.memory import desde_sitio

SITIO = "https://ejemplo.org"
CLABE = "012180001234567891"
RFC = "CEJ010203AB1"

PORTADA = f"""<!doctype html><html><head><title>Colectivo Ejemplo</title></head><body>
<p>Colectivo Ejemplo trabaja por la energía comunitaria.</p>
<p>Donativos a la CLABE {CLABE}. RFC {RFC}.</p>
<a href="/nosotros">Quiénes somos</a> <a href="/contacto">Contacto</a>
<a href="/galeria.jpg">Galería</a> <a href="https://otro.org/nosotros">Aliado</a>
</body></html>"""
NOSOTROS = "<html><body><p>Nuestra misión es llevar energía limpia a comunidades rurales.</p></body></html>"


def _cliente(robots: str = "User-agent: *\nDisallow: /contacto\n") -> httpx.Client:
    paginas = {"/robots.txt": robots, "/": PORTADA, "/nosotros": NOSOTROS, "/contacto": "<p>Tel 55 1234 5678</p>"}

    def responder(peticion: httpx.Request) -> httpx.Response:
        ruta = peticion.url.path or "/"
        if ruta not in paginas:
            return httpx.Response(404)
        tipo = "text/plain" if ruta.endswith(".txt") else "text/html; charset=utf-8"
        return httpx.Response(200, text=paginas[ruta], headers={"content-type": tipo})

    return httpx.Client(transport=httpx.MockTransport(responder))


def invocar(*args: str, entrada: str | None = None):
    return CliRunner().invoke(cli, list(args), input=entrada)


def _texto_docx(ruta) -> str:
    return "\n".join(p.text for p in Document(str(ruta)).paragraphs)


# --- Sitio web -----------------------------------------------------------------------


def test_rastrear_respeta_robots_mismo_dominio_y_quita_datos_sensibles():
    rastreo = desde_sitio.rastrear(SITIO, cliente=_cliente())
    assert [p.url for p in rastreo.paginas] == [SITIO, f"{SITIO}/nosotros"]
    assert any("contacto" in o and "robots.txt" in o for o in rastreo.omitidas)
    portada = rastreo.paginas[0]
    assert portada.titulo == "Colectivo Ejemplo"
    assert CLABE not in portada.texto and RFC not in portada.texto and DATO_BANCARIO_OMITIDO in portada.texto


def test_rastrear_falla_si_robots_prohibe_la_portada():
    with pytest.raises(ErrorLectura, match="robots.txt"):
        desde_sitio.rastrear(SITIO, cliente=_cliente("User-agent: *\nDisallow: /\n"))


def test_enlaces_utiles_ordena_por_pista_y_descarta_archivos_y_otros_dominios():
    enlaces = desde_sitio.enlaces_utiles(PORTADA, SITIO)
    assert enlaces == [f"{SITIO}/nosotros", f"{SITIO}/contacto"]


def test_interpretar_sin_cita_comprobable_deja_pendiente_y_filtra_mecanismos():
    paginas = desde_sitio.rastrear(SITIO, cliente=_cliente()).paginas
    respuesta = {
        "nombre": {"valor": "Colectivo Ejemplo", "evidencia": "Colectivo Ejemplo trabaja", "url": SITIO},
        "figura": {"valor": "Fundación", "evidencia": "una cita que no está", "url": SITIO},
        "mision": {"valor": "Energía limpia", "evidencia": "energía limpia a comunidades", "url": SITIO},
        "mecanismos_cobro": {"valor": ["Transferencia bancaria", "Criptomonedas"],
                             "evidencia": "Donativos a la CLABE", "url": SITIO},
    }
    propuestas = desde_sitio.interpretar(respuesta, paginas, desde_sitio.campos_del_sitio())
    assert propuestas["nombre"].valor == "Colectivo Ejemplo"
    assert propuestas["figura"].pendiente
    assert propuestas["mision"].pendiente  # la cita está en /nosotros, no en la portada
    assert propuestas["mecanismos_cobro"].valor == ["Transferencia bancaria"]
    assert not set(desde_sitio.CAMPOS_PRIVADOS) & set(propuestas)


def test_configurar_con_sitio_acepta_corrige_y_guarda_origenes(entorno_aislado, monkeypatch, respuestas_modelo):
    monkeypatch.setattr(desde_sitio, "nuevo_cliente", _cliente)
    respuestas_modelo.append(json.dumps({
        "nombre": {"valor": "Colectivo Ejemplo", "evidencia": "Colectivo Ejemplo trabaja", "url": SITIO},
        "figura": {"valor": "Fundación", "evidencia": "no aparece en ninguna página", "url": SITIO},
        "mision": {"valor": "Llevar energía limpia a comunidades rurales",
                   "evidencia": "Nuestra misión es llevar energía limpia", "url": f"{SITIO}/nosotros"},
        "mecanismos_cobro": {"valor": ["Transferencia bancaria"], "evidencia": "Donativos a la CLABE", "url": SITIO},
    }))
    # nombre, siglas, figura, misión; 13 campos sin cambio (incluye los privados); mecanismos; confirmar.
    entrada = "sí\n\nAsociación Civil\nsí\n" + "\n" * 13 + "sí\n\n"
    res = invocar("configurar", "--sitio", SITIO, entrada=entrada)
    assert res.exit_code == 0, res.output
    assert "Pendiente" in res.output and CLABE not in res.output

    perfil = entorno_aislado / "profile.yaml"
    datos = yaml.safe_load(perfil.read_text(encoding="utf-8"))
    assert datos["name"] == "Colectivo Ejemplo"
    assert datos["legal_structure"] == "Asociación Civil"
    assert datos["mission"] == "Llevar energía limpia a comunidades rurales"
    assert datos["procuracion"]["mecanismos_cobro"] == ["Transferencia bancaria"]
    assert datos["procuracion"]["origenes"] == {
        "nombre": SITIO, "figura": "Dato del usuario", "mision": f"{SITIO}/nosotros", "mecanismos_cobro": SITIO,
    }
    assert CLABE not in perfil.read_text(encoding="utf-8")
    assert list(entorno_aislado.glob("profile.yaml.bak-*"))


def test_configurar_manual_consulta_el_sitio_solo_para_esa_pregunta(entorno_aislado, monkeypatch, respuestas_modelo):
    monkeypatch.setattr(desde_sitio, "nuevo_cliente", _cliente)
    respuestas_modelo.append(json.dumps({
        "mision": {"valor": "Energía limpia para comunidades rurales",
                   "evidencia": "llevar energía limpia a comunidades rurales", "url": f"{SITIO}/nosotros"},
    }))
    # nombre, siglas, figura; misión «está en nuestro sitio», URL y «sí»; el resto sin cambio y confirmar.
    entrada = "Colectivo Ejemplo\n\n\nestá en nuestro sitio\n" + SITIO + "\nsí\n" + "\n" * 16
    res = invocar("configurar", entrada=entrada)
    assert res.exit_code == 0, res.output
    datos = yaml.safe_load((entorno_aislado / "profile.yaml").read_text(encoding="utf-8"))
    assert datos["mission"] == "Energía limpia para comunidades rurales"
    assert datos["procuracion"]["origenes"] == {"mision": f"{SITIO}/nosotros"}
    assert not respuestas_modelo


def test_configurar_no_admite_desde_y_sitio(entorno_aislado, tmp_path):
    archivo = tmp_path / "r.yaml"
    archivo.write_text("nombre: X\n", encoding="utf-8")
    res = invocar("configurar", "--desde", str(archivo), "--sitio", SITIO)
    assert res.exit_code == 1 and "no ambos" in res.output


# --- Reunión -------------------------------------------------------------------------


def test_validar_baja_etiquetas_sin_evidencia():
    pagina = desde_sitio.Pagina(url=SITIO, titulo="", texto="Financiamos proyectos de energía comunitaria.")
    prep = reunion.Preparacion.model_validate({
        "perfil_publico": [
            {"texto": "Financia energía", "etiqueta": "DATO PÚBLICO VERIFICADO", "fuente": SITIO,
             "cita": "proyectos de energía comunitaria"},
            {"texto": "Da 5 millones", "etiqueta": "DATO PÚBLICO VERIFICADO", "fuente": SITIO, "cita": "5 millones"},
        ],
        "puntos_en_comun": [{"texto": "Energía comunitaria", "etiqueta": "[dato institucional]"},
                            {"texto": "Misma visión", "etiqueta": "DATO INSTITUCIONAL"}],
        "riesgos": [{"texto": f"Piden depósito a {CLABE}", "etiqueta": "inventada"}],
    })
    validada = reunion.validar(prep, [pagina], "Ejes: Energía comunitaria")
    assert [d.texto for d in validada.perfil_publico] == ["Financia energía"]
    assert [d.etiqueta for d in validada.puntos_en_comun] == ["DATO INSTITUCIONAL", "INFERENCIA ESTRATÉGICA"]
    assert validada.riesgos[0].etiqueta == "NO VERIFICADO" and CLABE not in validada.riesgos[0].texto


def test_cli_reunion_preparar_con_word(entorno_aislado, respuestas_modelo):
    respuestas_modelo.append(json.dumps({
        "perfil_publico": [{"texto": "Financia energía", "etiqueta": "DATO PÚBLICO VERIFICADO",
                            "fuente": SITIO, "cita": "energía"}],
        "puntos_en_comun": [{"texto": "Ambos trabajan energía comunitaria", "etiqueta": "DATO INSTITUCIONAL"}],
        "riesgos": [{"texto": f"Piden la cuenta {CLABE}", "etiqueta": "x"}],
        "agenda": ["Presentación", "Prioridades del fondo"],
        "preguntas": ["¿Cuál es el calendario?"],
        "siguiente_paso": "Enviar nota conceptual",
    }))
    res = invocar("reunion", "preparar", "Fundación Ejemplo", "--objetivo", "Conocer el fondo", "--word", "--si")
    assert res.exit_code == 0, res.output
    assert "INFERENCIA ESTRATÉGICA" in res.output and "No se leyó un sitio público" in res.output
    assert "Financia energía" not in res.output and CLABE not in res.output
    docx = list((entorno_aislado / "salidas").rglob("*preparacion_reunion.docx"))
    assert len(docx) == 1
    texto = _texto_docx(docx[0])
    assert "Cooperativa de Prueba" in "\n".join(c.text for t in Document(str(docx[0])).tables for f in t.rows
                                               for c in f.cells)
    assert "Conocer el fondo" in texto and "Borrador" not in texto


# --- Redacción -----------------------------------------------------------------------


def test_cli_redactar_marca_cifras_sin_respaldo_y_guarda_word(entorno_aislado, respuestas_modelo):
    respuestas_modelo.append(
        "# Carta de intención\n\n**Techos solares**\n\nSolicitamos $500,000 MXN antes del 15 de diciembre "
        "de 2026.\n\n## Cierre\n\n- Quedamos atentos."
    )
    res = invocar("redactar", "carta-intencion", "--indicaciones", "Destacar techos solares", "--word", "--si")
    assert res.exit_code == 0, res.output
    assert "MONTO POR DEFINIR" in res.output and "Revisión automática" in res.output
    assert "Borrador: requiere revisión humana." in res.output
    docx = list((entorno_aislado / "salidas").rglob("*carta_intencion.docx"))
    assert len(docx) == 1
    texto = _texto_docx(docx[0])
    assert "500,000" not in texto and "MONTO POR DEFINIR" in texto and "Borrador" not in texto


def test_cli_redactar_expediente_inexistente(entorno_aislado):
    res = invocar("redactar", "propuesta", "--expediente", "EXP-2026-9999")
    assert res.exit_code == 1 and "No existe el expediente" in res.output


def test_desde_markdown_es_inverso_de_a_markdown():
    plano = DocumentoPlano(tipo="nota", titulo="Nota", subtitulo="Proyecto X", bloques=[
        Titulo("Objetivo"), Parrafo("Llevar energía."), Lista(["Uno", "Dos"]),
        Titulo("Detalle", 2), Tabla(["Dato", "Valor"], [["Monto", "MONTO POR DEFINIR"]]),
    ])
    assert a_markdown(desde_markdown.a_plano(a_markdown(plano), "nota")) == a_markdown(plano)


def test_desde_markdown_quita_enfasis_y_titulo_por_defecto():
    plano = desde_markdown.a_plano("Texto **importante** y *suave*.\n\n1. Paso uno", "correo_seguimiento")
    assert plano.titulo == "Correo seguimiento"
    assert isinstance(plano.bloques[0], Parrafo) and plano.bloques[0].texto == "Texto importante y suave."
    assert isinstance(plano.bloques[1], Lista) and plano.bloques[1].elementos == ["Paso uno"]


# --- Documentos nuevos -------------------------------------------------------------------


def _abrir_expediente() -> str:
    invocar("expedientes", "abrir", "Fundación Ejemplo", "--tipo", "Convocatoria", "--objetivo", "Techos solares")
    return cases.listar()[0].folio


def test_cli_carta_intencion_y_nota_conceptual(entorno_aislado):
    folio = _abrir_expediente()
    contenido = entorno_aislado / "carta.yaml"
    contenido.write_text(yaml.safe_dump({"monto_solicitado": 250000, "firmante": "Dirección General"},
                                        allow_unicode=True), encoding="utf-8")
    res = invocar("documento", "carta-intencion", "--expediente", folio, "--contenido", str(contenido), "--si")
    assert res.exit_code == 0, res.output
    carta = list((entorno_aislado / "salidas").rglob("*_carta_intencion.docx"))
    assert len(carta) == 1
    texto = _texto_docx(carta[0])
    assert "Fundación Ejemplo" in texto and "250,000" in texto and "Dirección General" in texto

    res = invocar("documento", "nota-conceptual", "--expediente", folio, "--si")
    assert res.exit_code == 0, res.output
    assert "Datos pendientes" in res.output
    assert len(list((entorno_aislado / "salidas").rglob("*_nota_conceptual.docx"))) == 1


def test_cli_carta_intencion_contenido_invalido(entorno_aislado):
    contenido = entorno_aislado / "carta.yaml"
    contenido.write_text("monto_solicitado: -5\n", encoding="utf-8")
    res = invocar("documento", "carta-intencion", "--contenido", str(contenido), "--si")
    assert res.exit_code == 1


def _guardar_evaluacion(nombre: str, dias: int) -> int:
    hoy = date.today()
    datos = {
        "convocatoria": nombre, "financiador": "Fundación Ejemplo", "elegibilidad": "confirmada",
        "vigencia": "vigente", "fecha_cierre": (hoy + timedelta(days=dias)).isoformat(), "evidencia_plazo": "Bases.",
        "criterios": {c: {"puntos": p, "evidencia": "Bases."} for c, p in PESOS.items()},
    }
    return matrix.guardar(matrix.evaluar(datos, hoy=hoy, rango_presupuesto_definido=True))


def test_cli_comparar_convocatorias_y_reporte_evaluacion(entorno_aislado):
    lejana = _guardar_evaluacion("Fondo Lejano", 90)
    cercana = _guardar_evaluacion("Fondo Cercano", 45)
    res = invocar("comparar-convocatorias", str(lejana), str(cercana))
    assert res.exit_code == 0, res.output
    assert "Qué atender primero" in res.output and f"#{cercana} Fondo Cercano" in res.output

    res = invocar("documento", "reporte-evaluacion", "--si")
    assert res.exit_code == 0, res.output
    assert len(list((entorno_aislado / "salidas").rglob("*_reporte_evaluacion.docx"))) == 1

    assert invocar("comparar-convocatorias", str(lejana)).exit_code == 1


# --- Análisis ------------------------------------------------------------------------


def test_cli_fecha(entorno_aislado):
    res = invocar("fecha", "--cierre", "2099-01-01")
    assert res.exit_code == 0 and "Hoy es" in res.output and "Faltan" in res.output
    res = invocar("fecha", "--cierre", "pronto")
    assert res.exit_code == 1 and "no es una fecha" in res.output


def test_cli_consultar_sin_modelo(entorno_aislado):
    carpeta = entorno_aislado / "conocimiento"
    carpeta.mkdir()
    (carpeta / "politicas.md").write_text("# Donativos\n\nSolo aceptamos donativos en especie.", encoding="utf-8")
    res = invocar("consultar", "¿Qué donativos aceptan en especie?")
    assert res.exit_code == 0 and "conocimiento/politicas.md" in res.output
    res = invocar("consultar", "astronomía cuántica")
    assert "No consta" in res.output


def test_cli_consultar_redactar_revisa_la_respuesta(entorno_aislado, respuestas_modelo):
    respuestas_modelo.append("Atendemos a 3,000 familias con un fondo de $2,000,000 MXN [1].")
    res = invocar("consultar", "bienestar común", "--redactar")
    assert res.exit_code == 0, res.output
    assert "MONTO POR DEFINIR" in res.output and "Fuentes:" in res.output


def test_cli_revisar_archivo(entorno_aislado, respuestas_modelo):
    bases = entorno_aislado / "bases.md"
    bases.write_text(
        "## Requisitos\n\nSer asociación civil con dos años de constitución.\n\n"
        f"## Pagos\n\nLos apoyos se depositan a la CLABE {CLABE}.",
        encoding="utf-8",
    )
    res = invocar("revisar", str(bases), "--sin-modelo", "--pregunta", "¿Qué requisitos piden?")
    assert res.exit_code == 0 and "asociación civil" in res.output
    res = invocar("revisar", str(bases), "--sin-modelo", "--pregunta", "CLABE depósito")
    assert CLABE not in res.output
    assert invocar("revisar", str(bases), "--sin-modelo").exit_code == 1
    assert invocar("revisar", "no-existe.pdf").exit_code == 1

    respuestas_modelo.append("Piden ser asociación civil y otorgan $1,000,000 MXN.")
    res = invocar("revisar", str(bases))
    assert res.exit_code == 0 and "MONTO POR DEFINIR" in res.output and "asociación civil" in res.output
