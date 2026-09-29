"""Lectura de archivos y URL sin red: MockTransport y un PDF armado a mano."""
import httpx
import pytest
from docx import Document

from coopexecutive import lectura
from coopexecutive.lectura import ErrorLectura, leer_archivo, leer_url, texto_de_html


def pdf_con_texto(texto: str) -> bytes:
    """PDF mínimo de una página con el texto en Helvetica."""
    flujo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(flujo) + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida, posiciones = b"%PDF-1.4\n", []
    for n, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += b"%d 0 obj\n" % n + cuerpo + b"\nendobj\n"
    xref = len(salida)
    salida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    salida += b"".join(b"%010d 00000 n \n" % p for p in posiciones)
    salida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objetos) + 1, xref)
    return salida


# --- Archivos --------------------------------------------------------------------


def test_texto_de_html_quita_scripts_y_separa_bloques():
    texto = texto_de_html("<html><script>x=1</script><h1>Bases</h1><p>Monto &amp; plazo</p></html>")
    assert texto.splitlines() == ["Bases", "Monto & plazo"]


def test_leer_archivo_txt_html_y_cp1252(tmp_path):
    (tmp_path / "a.txt").write_text("Convocatoria única", encoding="utf-8")
    (tmp_path / "b.html").write_text("<p>Hola</p><style>p{}</style>", encoding="utf-8")
    (tmp_path / "c.md").write_bytes("Año".encode("cp1252"))
    assert leer_archivo(tmp_path / "a.txt") == "Convocatoria única"
    assert leer_archivo(tmp_path / "b.html") == "Hola"
    assert leer_archivo(tmp_path / "c.md") == "Año"


def test_leer_archivo_docx(tmp_path):
    documento = Document()
    documento.add_paragraph("Requisitos de elegibilidad")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text, tabla.cell(0, 1).text = "Monto", "Por definir"
    documento.save(tmp_path / "bases.docx")
    texto = leer_archivo(tmp_path / "bases.docx")
    assert "Requisitos de elegibilidad" in texto
    assert "Monto | Por definir" in texto


def test_leer_archivo_pdf(tmp_path):
    pytest.importorskip("pypdf")
    (tmp_path / "bases.pdf").write_bytes(pdf_con_texto("Cierre de la convocatoria"))
    assert "Cierre de la convocatoria" in leer_archivo(tmp_path / "bases.pdf")


def test_leer_archivo_errores(tmp_path):
    with pytest.raises(ErrorLectura, match="No se pudo abrir"):
        leer_archivo(tmp_path / "no-existe.txt")
    (tmp_path / "x.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ErrorLectura, match=r"\.doc no se admite"):
        leer_archivo(tmp_path / "x.doc")
    (tmp_path / "roto.pdf").write_bytes(b"%PDF-1.4\nbasura")
    with pytest.raises(ErrorLectura):
        leer_archivo(tmp_path / "roto.pdf")


# --- URL -------------------------------------------------------------------------


def cliente(manejador) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(manejador))


def test_leer_url_html_y_pdf():
    pytest.importorskip("pypdf")

    def manejador(peticion: httpx.Request) -> httpx.Response:
        if peticion.url.path == "/bases.pdf":
            return httpx.Response(200, content=pdf_con_texto("Bases en PDF"),
                                  headers={"content-type": "application/pdf"})
        return httpx.Response(200, text="<html><body><p>Bases web</p></body></html>",
                              headers={"content-type": "text/html; charset=utf-8"})

    c = cliente(manejador)
    assert leer_url("https://fondo.org/", cliente=c) == "Bases web"
    assert "Bases en PDF" in leer_url("https://fondo.org/bases.pdf", cliente=c)


def test_leer_url_corta_en_el_tope():
    c = cliente(lambda p: httpx.Response(200, content=b"a" * 2000, headers={"content-type": "text/plain"}))
    with pytest.raises(ErrorLectura, match="No se pudo descargar"):
        leer_url("https://fondo.org/grande", cliente=c, tope=1000)


def test_leer_url_error_http_y_tipo_no_admitido():
    c = cliente(lambda p: httpx.Response(404))
    with pytest.raises(ErrorLectura, match="HTTP 404"):
        leer_url("https://fondo.org/x", cliente=c)
    c = cliente(lambda p: httpx.Response(200, content=b"\x89PNG", headers={"content-type": "image/png"}))
    with pytest.raises(ErrorLectura, match="no admitido"):
        leer_url("https://fondo.org/logo.png", cliente=c)


def test_bloquea_direcciones_internas_incluso_tras_redireccion(monkeypatch):
    resoluciones = {"fondo.org": "93.184.216.34", "interno.local": "10.0.0.5"}

    def resolver(host, *_):
        return [(None, None, None, None, (resoluciones[host], 0))]

    monkeypatch.setattr(lectura.socket, "getaddrinfo", resolver)

    def manejador(peticion: httpx.Request) -> httpx.Response:
        if peticion.url.host == "fondo.org":
            return httpx.Response(302, headers={"location": "http://interno.local/admin"})
        return httpx.Response(200, text="secreto", headers={"content-type": "text/plain"})

    with pytest.raises(ErrorLectura, match="dirección interna"):
        leer_url("https://fondo.org/", cliente=cliente(manejador), bloquear_privadas=True)
    # Sin el bloqueo (uso local desde la CLI), la redirección se sigue.
    assert leer_url("https://fondo.org/", cliente=cliente(manejador)) == "secreto"


def test_rechaza_esquemas_no_web():
    with pytest.raises(ErrorLectura, match="URL no admitida"):
        leer_url("file:///etc/passwd", cliente=cliente(lambda p: httpx.Response(200)), bloquear_privadas=True)
