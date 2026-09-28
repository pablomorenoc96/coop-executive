import hashlib
from datetime import date, timedelta

import pytest
import yaml
from click.testing import CliRunner
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from coopexecutive import documents
from coopexecutive.cli import cli
from coopexecutive.crm import cases
from coopexecutive.documents import generador
from coopexecutive.documents.contenido import DocumentoPlano, Parrafo, Tabla, Titulo, a_markdown
from coopexecutive.documents.formato import a_docx, ancho_util
from coopexecutive.grant_tools import matrix
from coopexecutive.grant_tools.matrix import PESOS
from coopexecutive.memory.company_profile import CoopProfile, Membrete

HOY = date(2026, 9, 27)

PERFIL = CoopProfile(
    name="Cooperativa Ejemplo",
    legal_structure="S.C. de R.L.",
    mission="Impulsar la energía comunitaria.",
    target_communities=["Comunidades rurales"],
    procuracion={"territorio": "Región Centro", "programas": ["Techos solares"], "alianzas": ["Red Ejemplo"]},
)


def _expediente(**extra) -> cases.Expediente:
    datos = {"folio": "EXP-2026-0001", "tipo": "Convocatoria", "entidad": "Fundación Ejemplo",
             "titulo": "Convocatoria - Fundación Ejemplo", **extra}
    return cases.Expediente(**datos)


def _resultado(**extra) -> matrix.ResultadoMatriz:
    datos = {
        "convocatoria": "Fondo de Energía Comunitaria",
        "financiador": "Fundación Ejemplo",
        "elegibilidad": "confirmada",
        "vigencia": "vigente",
        "fecha_cierre": (HOY + timedelta(days=40)).isoformat(),
        "evidencia_plazo": "Bases, numeral 5.",
        "criterios": {c: {"puntos": p, "evidencia": f"Bases sobre {c}."} for c, p in PESOS.items()},
        **extra,
    }
    return matrix.evaluar(datos, hoy=HOY, rango_presupuesto_definido=True)


def _texto(plano: DocumentoPlano) -> str:
    return "\n".join(plano.textos())


# --- Contenido ---------------------------------------------------------------------


def test_solicitud_sin_contenido_marca_pendientes_y_no_lleva_datos_internos():
    plano = generador.solicitud(PERFIL, _expediente(), None, HOY)
    texto = _texto(plano)
    faltantes = plano.pendientes()
    assert "[PENDIENTE: nombre del proyecto]" in faltantes
    assert "[PENDIENTE: justificación del proyecto]" in faltantes
    assert "MONTO POR DEFINIR" in faltantes and "COSTO POR COTIZAR" in faltantes
    assert "27 de septiembre de 2026" in texto
    assert "Techos solares" in texto and "Red Ejemplo" in texto
    # Documento para terceros: sin folio ni huellas.
    assert "EXP-" not in texto and "SHA-256" not in texto


def test_solicitud_con_contenido_suma_solo_montos_completos_de_una_moneda():
    contenido = generador.ContenidoSolicitud(
        proyecto="Techos solares comunitarios",
        presupuesto=[{"concepto": "Paneles", "monto": 100000}, {"concepto": "Instalación", "monto": 25000}],
        moneda="MXN",
    )
    plano = generador.solicitud(PERFIL, _expediente(tipo="Beca"), contenido, HOY)
    assert plano.titulo == "Solicitud de beca"
    tabla = next(b for b in plano.bloques if isinstance(b, Tabla) and b.encabezados == ["Concepto", "Monto"])
    assert tabla.filas[-1] == ["Total", "125,000.00 MXN"]

    contenido.presupuesto.append(generador.Partida(concepto="Capacitación", monto=None))
    tabla = generador._tabla_presupuesto(contenido, "MXN")
    assert ["Capacitación", "COSTO POR COTIZAR"] in tabla.filas
    assert tabla.filas[-1] == ["Total", "MONTO POR DEFINIR"]

    mezcla = generador.ContenidoSolicitud(presupuesto=[
        {"concepto": "A", "monto": 10, "moneda": "USD"}, {"concepto": "B", "monto": 10, "moneda": "MXN"},
    ])
    assert generador._tabla_presupuesto(mezcla, "MXN").filas[-1] == ["Total", "MONTO POR DEFINIR"]


def test_contenido_con_moneda_invalida_se_rechaza():
    with pytest.raises(ValueError):
        generador.ContenidoSolicitud(moneda="PESOS")


def test_documento_institucional_omite_opcionales_vacios():
    plano = generador.documento_institucional(PERFIL, HOY)
    titulos = [b.texto for b in plano.bloques if isinstance(b, Titulo)]
    assert "Misión" in titulos and "Programas" in titulos
    assert "Visión" not in titulos and "Cómo apoyar" not in titulos
    assert "Gobierno" in titulos  # las cooperativas tienen asamblea
    assert "[PENDIENTE: métricas de impacto verificables]" in plano.pendientes()


def test_ficha_no_escribe_tensiones_internas():
    resultado = _resultado(tension="pendiente", tension_detalle="Diferencias con un aliado")
    assert resultado.decision == "CONDICIONAL" and "tensión" in resultado.motivo
    plano = generador.ficha_oportunidad(PERFIL, resultado, _expediente())
    texto = _texto(plano)
    assert "tensión" not in texto.lower() and "Diferencias con un aliado" not in texto
    assert "Decisión: Condicional" in texto
    assert resultado.hash in texto and "EXP-2026-0001" in texto


def test_vista_previa_en_markdown():
    plano = DocumentoPlano("prueba", "Título", "Subtítulo", [
        Titulo("Sección"), Parrafo("Texto"), Tabla(["A", "B"], [["1", "x|y"]]),
    ])
    md = a_markdown(plano)
    assert md.startswith("# Título") and "## Sección" in md and "| 1 | x/y |" in md


# --- Formato Word --------------------------------------------------------------------


def _imagen(ruta, tamano=(1600, 180)):
    PIL = pytest.importorskip("PIL")
    from PIL import Image

    Image.new("RGB", tamano, (40, 40, 40)).save(ruta)
    return ruta


def test_formato_institucional_y_membrete(tmp_path):
    _imagen(tmp_path / "membrete.png")
    membrete = Membrete(imagen="membrete.png", pie="Calle Ejemplo 1, Región Centro", fuente="Arial", tamano=11)
    doc = a_docx(generador.solicitud(PERFIL, _expediente(), None, HOY), membrete, tmp_path, autor=PERFIL.name)

    normal = doc.styles["Normal"]
    assert normal.font.name == "Arial" and normal.font.size == Pt(11)
    assert normal.paragraph_format.line_spacing == 1.15
    assert normal.paragraph_format.space_after == Pt(6)
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    seccion = doc.sections[0]
    corrida = seccion.header.paragraphs[0].runs[0]._r
    assert len(corrida.xpath(".//pic:pic")) == 1
    # Imagen ancha: ocupa el ancho útil.
    assert int(corrida.xpath(".//wp:extent")[0].get("cx")) == ancho_util(doc)
    assert "Calle Ejemplo 1" in seccion.footer.paragraphs[0].text
    assert doc.core_properties.author == PERFIL.name
    for tabla in doc.tables:
        assert abs(sum(c.width for c in tabla.columns) - ancho_util(doc)) < 10


def test_ficha_filas_enteras_y_huella_sin_justificar(tmp_path):
    resultado = _resultado()
    doc = a_docx(generador.ficha_oportunidad(PERFIL, resultado, _expediente()), Membrete(), tmp_path)

    for tabla in doc.tables:
        for fila in tabla.rows:
            assert fila._tr.trPr.xpath("./w:cantSplit")
    huella = next(p for p in doc.paragraphs if resultado.hash in p.text)
    assert huella.alignment == WD_ALIGN_PARAGRAPH.LEFT
    # «Pendiente» cabe en la columna de puntos sin partirse.
    puntos = doc.tables[-1].columns[2].width
    assert puntos > ancho_util(doc) * 0.14


def test_logo_cuadrado_se_limita_por_altura(tmp_path):
    from docx.shared import Cm

    _imagen(tmp_path / "logo.png", (400, 400))
    doc = a_docx(DocumentoPlano("x", "X"), Membrete(imagen="logo.png"), tmp_path)
    forma = doc.sections[0].header.paragraphs[0].runs[0]._r.xpath(".//wp:extent")[0]
    assert int(forma.get("cy")) == Cm(2.5)


def test_imagen_de_membrete_inexistente(tmp_path):
    with pytest.raises(FileNotFoundError, match="no_existe.png"):
        a_docx(DocumentoPlano("x", "X"), Membrete(imagen="no_existe.png"), tmp_path)


def test_guardar_no_sobrescribe_y_registra_huella(entorno_aislado):
    plano = generador.documento_institucional(PERFIL, HOY)
    salidas = entorno_aislado / "salidas"
    primero = documents.guardar(plano, Membrete(), entorno_aislado, salidas, HOY, expediente_folio="EXP-2026-0001")
    segundo = documents.guardar(plano, Membrete(), entorno_aislado, salidas, HOY, expediente_folio="EXP-2026-0001")
    assert primero.ruta != segundo.ruta
    assert segundo.ruta.endswith("2026-09-27_documento_institucional-2.docx")
    assert "2026" in primero.ruta and "EXP-2026-0001" in primero.ruta
    with open(primero.ruta, "rb") as archivo:
        assert hashlib.sha256(archivo.read()).hexdigest() == primero.sha256
    assert [d.id for d in documents.documentos_de("exp-2026-0001")] == [primero.id, segundo.id]
    Document(primero.ruta)  # el archivo abre


# --- CLI ------------------------------------------------------------------------------------


def invocar(*args: str, entrada: str | None = None):
    return CliRunner().invoke(cli, list(args), input=entrada)


def test_cli_solicitud_con_confirmacion(entorno_aislado):
    invocar("expedientes", "abrir", "Fundación Ejemplo", "--tipo", "Convocatoria", "--objetivo", "Techos solares")
    folio = cases.listar()[0].folio
    contenido = entorno_aislado / "contenido.yaml"
    contenido.write_text(yaml.safe_dump({"proyecto": "Techos solares comunitarios"}, allow_unicode=True),
                         encoding="utf-8")

    res = invocar("documento", "solicitud", "--expediente", folio, "--contenido", str(contenido), entrada="n\n")
    assert res.exit_code == 0, res.output
    assert "Vista previa" in res.output and "Datos pendientes" in res.output
    assert "No se generó" in res.output
    assert not (entorno_aislado / "salidas").exists()

    res = invocar("documento", "solicitud", "--expediente", folio, "--contenido", str(contenido), "--si")
    assert res.exit_code == 0, res.output
    assert "Documento guardado" in res.output
    assert len(list((entorno_aislado / "salidas").rglob("*_solicitud.docx"))) == 1

    res = invocar("expedientes", "ver", folio)
    assert "Documentos" in res.output and "solicitud" in res.output


def test_cli_ficha_desde_expediente(entorno_aislado):
    invocar("expedientes", "abrir", "Fundación Ejemplo", "--tipo", "Convocatoria")
    folio = cases.listar()[0].folio
    res = invocar("documento", "ficha", "--expediente", folio, "--si")
    assert res.exit_code == 1 and "no tiene evaluaciones" in res.output

    id_evaluacion = matrix.guardar(_resultado(), folio)
    res = invocar("documento", "ficha", "--evaluacion", str(id_evaluacion), "--si")
    assert res.exit_code == 0, res.output
    assert f"Registrado en el expediente {folio}" in res.output
    assert documents.documentos_de(folio)[0].evaluacion_id == id_evaluacion


def test_cli_ficha_sin_argumentos_y_evaluacion_inexistente(entorno_aislado):
    assert "Indique --expediente o --evaluacion" in invocar("documento", "ficha").output
    res = invocar("documento", "ficha", "--evaluacion", "99")
    assert res.exit_code == 1 and "No existe la evaluación 99" in res.output


def test_cli_institucional_sin_expediente(entorno_aislado):
    res = invocar("documento", "institucional", "--si")
    assert res.exit_code == 0, res.output
    assert len(list((entorno_aislado / "salidas").rglob("general/*_documento_institucional.docx"))) == 1
