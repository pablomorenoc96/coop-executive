"""Convierte un ``DocumentoPlano`` en Word con el membrete de la organización.

Formato: carta con márgenes de 2.5 cm; la fuente y el tamaño del membrete (Arial 11
por omisión); interlineado 1.15; 6 pt después de cada párrafo; texto justificado;
tablas al ancho de la página con la fila de encabezado repetida en cada página y
filas que no se parten entre páginas.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as ImagenDocx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from coopexecutive.documents.contenido import DocumentoPlano, Lista, Parrafo, Tabla, Titulo
from coopexecutive.memory.company_profile import Membrete
from coopexecutive.utils.rutas import resolver_archivo

ANCHO_CARTA = Cm(21.59)
ALTO_CARTA = Cm(27.94)
MARGEN = Cm(2.5)
ALTO_MAXIMO_MEMBRETE = Cm(2.5)
COLOR_TITULOS = RGBColor(0x1F, 0x1F, 0x1F)
SOMBRA_ENCABEZADO = "E7E6E6"


def _fijar_fuente(estilo, fuente: str) -> None:
    """Asigna la fuente a todos los alfabetos y quita la del tema, que tendría prioridad."""
    estilo.font.name = fuente
    rpr = estilo.element.get_or_add_rPr()
    fuentes = rpr.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        rpr.insert(0, fuentes)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fuentes.set(qn(atributo), fuente)
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fuentes.attrib.pop(qn(atributo), None)


def _estilos(doc, membrete: Membrete) -> None:
    normal = doc.styles["Normal"]
    _fijar_fuente(normal, membrete.fuente)
    normal.font.size = Pt(membrete.tamano)
    formato = normal.paragraph_format
    formato.line_spacing = 1.15
    formato.space_before = Pt(0)
    formato.space_after = Pt(6)
    formato.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for nombre, extra in (("Heading 1", 3), ("Heading 2", 1)):
        estilo = doc.styles[nombre]
        _fijar_fuente(estilo, membrete.fuente)
        estilo.font.size = Pt(membrete.tamano + extra)
        estilo.font.bold = True
        estilo.font.color.rgb = COLOR_TITULOS
        estilo.paragraph_format.space_before = Pt(12)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True
        estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lista = doc.styles["List Bullet"]
    _fijar_fuente(lista, membrete.fuente)
    lista.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lista.paragraph_format.space_after = Pt(3)


def _pagina(doc) -> None:
    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = ANCHO_CARTA, ALTO_CARTA
    seccion.left_margin = seccion.right_margin = MARGEN
    seccion.top_margin = seccion.bottom_margin = MARGEN
    seccion.header_distance = seccion.footer_distance = Cm(1.25)


def ancho_util(doc) -> Emu:
    seccion = doc.sections[0]
    return Emu(seccion.page_width - seccion.left_margin - seccion.right_margin)


def _numero_de_pagina(parrafo) -> None:
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    corrida = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "1"
    corrida.append(texto)
    campo.append(corrida)
    parrafo._p.append(campo)


def _membrete(doc, membrete: Membrete, carpeta_perfil: Path) -> None:
    seccion = doc.sections[0]
    imagen = resolver_archivo(membrete.imagen, carpeta_perfil)
    if imagen is not None:
        parrafo = seccion.header.paragraphs[0]
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # Separa el membrete del primer título de cada página.
        parrafo.paragraph_format.space_after = Pt(12)
        medidas = ImagenDocx.from_file(str(imagen))
        ancho = ancho_util(doc)
        # Un encabezado ancho ocupa el ancho útil; un logo cuadrado se limita por altura.
        if ancho * medidas.px_height / medidas.px_width > ALTO_MAXIMO_MEMBRETE:
            parrafo.add_run().add_picture(str(imagen), height=ALTO_MAXIMO_MEMBRETE)
        else:
            parrafo.add_run().add_picture(str(imagen), width=ancho)
    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if membrete.pie.strip():
        corrida = pie.add_run(membrete.pie.strip())
        corrida.font.size = Pt(max(membrete.tamano - 2, 7))
        pie = seccion.footer.add_paragraph()
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _numero_de_pagina(pie)
    for p in seccion.footer.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for corrida in p.runs:
            corrida.font.size = Pt(max(membrete.tamano - 2, 7))


def _sombrear(celda, color: str) -> None:
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(sombra)


def _marcar_fila(fila, etiqueta: str) -> None:
    marca = OxmlElement(etiqueta)
    marca.set(qn("w:val"), "true")
    fila._tr.get_or_add_trPr().append(marca)


def _texto_celda(celda, texto: str, tamano: float, negrita: bool = False) -> None:
    parrafo = celda.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    parrafo.paragraph_format.space_after = Pt(0)
    corrida = parrafo.add_run(texto)
    corrida.bold = negrita
    corrida.font.size = Pt(tamano)


def _tabla(doc, bloque: Tabla, membrete: Membrete) -> None:
    columnas = len(bloque.encabezados)
    tabla = doc.add_table(rows=1, cols=columnas)
    tabla.style = doc.styles["Table Grid"]
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    pesos = bloque.anchos if len(bloque.anchos) == columnas else [1.0] * columnas
    total = sum(pesos)
    anchos = [Emu(int(ancho_util(doc) * p / total)) for p in pesos]
    tamano = max(membrete.tamano - 1, 7)
    for celda, texto in zip(tabla.rows[0].cells, bloque.encabezados):
        _texto_celda(celda, texto, tamano, negrita=True)
        _sombrear(celda, SOMBRA_ENCABEZADO)
    _marcar_fila(tabla.rows[0], "w:tblHeader")
    for fila in bloque.filas:
        nueva = tabla.add_row()
        for celda, texto in zip(nueva.cells, fila):
            _texto_celda(celda, texto, tamano)
    for fila in tabla.rows:
        _marcar_fila(fila, "w:cantSplit")
    for columna, ancho in zip(tabla.columns, anchos):
        columna.width = ancho
        for celda in columna.cells:
            celda.width = ancho
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def a_docx(plano: DocumentoPlano, membrete: Membrete, carpeta_perfil: Path, autor: str = ""):
    """Devuelve el ``Document`` de python-docx listo para guardar."""
    doc = Document()
    _pagina(doc)
    _estilos(doc, membrete)
    _membrete(doc, membrete, carpeta_perfil)
    propiedades = doc.core_properties
    propiedades.title = plano.titulo
    propiedades.author = propiedades.last_modified_by = autor
    propiedades.comments = ""

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = titulo.add_run(plano.titulo)
    corrida.bold = True
    corrida.font.size = Pt(membrete.tamano + 5)
    if plano.subtitulo:
        subtitulo = doc.add_paragraph()
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        corrida = subtitulo.add_run(plano.subtitulo)
        corrida.bold = True
        corrida.font.size = Pt(membrete.tamano + 1)

    for bloque in plano.bloques:
        if isinstance(bloque, Titulo):
            doc.add_heading(bloque.texto, level=min(max(bloque.nivel, 1), 2))
        elif isinstance(bloque, Parrafo):
            parrafo = doc.add_paragraph(bloque.texto)
            if not bloque.justificado:
                parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif isinstance(bloque, Lista):
            for elemento in bloque.elementos:
                doc.add_paragraph(elemento, style="List Bullet")
        else:
            _tabla(doc, bloque, membrete)
    return doc
