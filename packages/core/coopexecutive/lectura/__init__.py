"""Lectura de bases y páginas: archivos locales (txt, md, html, docx, pdf) y URL.

Las descargas se hacen por partes y se cortan al llegar al tope de bytes. Cuando la
petición viene de fuera (MCP o la API local), ``bloquear_privadas`` impide que una
URL apunte a la red interna, también tras una redirección.
"""
from __future__ import annotations

import html
import io
import ipaddress
import re
import socket
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx

from coopexecutive import __version__

AGENTE = f"Mozilla/5.0 (compatible; CoopExecutive/{__version__}; +https://github.com/pablomorenoc96/coop-executive)"
TOPE_BYTES = 5_000_000
TIEMPO_ESPERA = 20.0
MAX_REDIRECCIONES = 5

__all__ = ["ErrorLectura", "descargar_bytes", "leer_archivo", "leer_url", "texto_de_html"]


class ErrorLectura(ValueError):
    """No se pudo obtener texto del origen indicado."""


# --- Texto -------------------------------------------------------------------------

_BLOQUE = re.compile(r"</?(p|div|br|li|tr|h[1-6]|section|article|table|ul|ol)\b[^>]*>", re.IGNORECASE)
_INVISIBLE = re.compile(r"<(script|style|noscript|template)\b.*?</\1\s*>", re.IGNORECASE | re.DOTALL)


def texto_de_html(fuente: str) -> str:
    """Texto legible de una página: sin scripts ni estilos y con un salto por bloque."""
    fuente = _INVISIBLE.sub(" ", fuente)
    fuente = re.sub(r"<!--.*?-->", " ", fuente, flags=re.DOTALL)
    fuente = _BLOQUE.sub("\n", fuente)
    fuente = html.unescape(re.sub(r"<[^>]+>", " ", fuente))
    lineas = (re.sub(r"[ \t\r\f\v]+", " ", linea).strip() for linea in fuente.splitlines())
    return "\n".join(linea for linea in lineas if linea)


def _texto_de_pdf(contenido: bytes) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise ErrorLectura(
            "Para leer PDF instale el complemento: pip install \"coopexecutive[pdf]\"."
        ) from exc
    try:
        lector = PdfReader(io.BytesIO(contenido))
        if lector.is_encrypted:
            raise ErrorLectura("El PDF está protegido con contraseña.")
        paginas = [pagina.extract_text() or "" for pagina in lector.pages]
    except PdfReadError as exc:
        raise ErrorLectura(f"El PDF no se pudo leer: {exc}") from exc
    texto = "\n\n".join(p.strip() for p in paginas if p.strip())
    if not texto:
        raise ErrorLectura("El PDF no tiene texto extraíble (puede ser una imagen escaneada).")
    return texto


def _texto_de_docx(contenido: bytes) -> str:
    from docx import Document

    documento = Document(io.BytesIO(contenido))
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    return "\n".join(partes)


def _decodificar(contenido: bytes) -> str:
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return contenido.decode("latin-1")


def leer_archivo(ruta: str | Path) -> str:
    """Texto de un archivo local según su extensión."""
    ruta = Path(ruta)
    try:
        contenido = ruta.read_bytes()
    except OSError as exc:
        raise ErrorLectura(f"No se pudo abrir {ruta.name}: {exc.strerror or exc}") from exc
    extension = ruta.suffix.lower()
    if extension == ".pdf" or contenido.startswith(b"%PDF-"):
        return _texto_de_pdf(contenido)
    if extension == ".docx":
        try:
            return _texto_de_docx(contenido)
        except Exception as exc:  # python-docx lanza varios tipos ante un archivo dañado
            raise ErrorLectura(f"{ruta.name} no es un Word válido.") from exc
    if extension == ".doc":
        raise ErrorLectura("El formato .doc no se admite; guárdelo como .docx o PDF.")
    texto = _decodificar(contenido)
    if extension in (".html", ".htm"):
        return texto_de_html(texto)
    if "\x00" in texto:
        raise ErrorLectura(f"{ruta.name} no es un archivo de texto.")
    return texto


# --- Red ---------------------------------------------------------------------------


def _verificar_destino(url: str) -> None:
    partes = urlparse(url)
    if partes.scheme not in ("http", "https") or not partes.hostname:
        raise ErrorLectura(f"URL no admitida: {url}")
    try:
        direcciones = {info[4][0] for info in socket.getaddrinfo(partes.hostname, None)}
    except socket.gaierror as exc:
        raise ErrorLectura(f"No se encontró el sitio {partes.hostname}.") from exc
    for direccion in direcciones:
        ip = ipaddress.ip_address(direccion.split("%")[0])
        if not ip.is_global:
            raise ErrorLectura(f"La URL apunta a una dirección interna ({partes.hostname}); no se consulta.")


def nuevo_cliente() -> httpx.Client:
    return httpx.Client(timeout=TIEMPO_ESPERA, follow_redirects=False, headers={"User-Agent": AGENTE})


def descargar_bytes(
    url: str,
    cliente: httpx.Client,
    *,
    tope: int = TOPE_BYTES,
    params: dict | None = None,
    bloquear_privadas: bool = False,
) -> tuple[bytes, str]:
    """(contenido, tipo) de la URL. Corta la descarga al pasar el tope.

    Las redirecciones se siguen aquí, una por una, para revisar cada destino.
    Lanza httpx.HTTPError si la red o el servidor fallan.
    """
    actual = url
    for _ in range(MAX_REDIRECCIONES + 1):
        if bloquear_privadas:
            _verificar_destino(actual)
        with cliente.stream("GET", actual, params=params) as respuesta:
            if respuesta.is_redirect and "location" in respuesta.headers:
                actual = urljoin(str(respuesta.url), respuesta.headers["location"])
                params = None
                continue
            respuesta.raise_for_status()
            declarado = respuesta.headers.get("content-length", "")
            if declarado.isdigit() and int(declarado) > tope:
                raise httpx.HTTPError(f"respuesta de más de {tope // 1_000_000} MB")
            partes, total = [], 0
            for parte in respuesta.iter_bytes():
                total += len(parte)
                if total > tope:
                    raise httpx.HTTPError(f"respuesta de más de {tope // 1_000_000} MB")
                partes.append(parte)
            return b"".join(partes), respuesta.headers.get("content-type", "").split(";")[0].strip().lower()
    raise httpx.HTTPError(f"demasiadas redirecciones desde {url}")


def leer_url(
    url: str,
    *,
    cliente: httpx.Client | None = None,
    tope: int = TOPE_BYTES,
    bloquear_privadas: bool = False,
) -> str:
    """Texto de una página web, un PDF o un archivo de texto publicado."""
    propio = cliente is None
    cliente = cliente or nuevo_cliente()
    try:
        contenido, tipo = descargar_bytes(url, cliente, tope=tope, bloquear_privadas=bloquear_privadas)
    except httpx.HTTPStatusError as exc:
        raise ErrorLectura(f"El sitio respondió HTTP {exc.response.status_code}.") from exc
    except httpx.HTTPError as exc:
        raise ErrorLectura(f"No se pudo descargar {url}: {exc or type(exc).__name__}") from exc
    finally:
        if propio:
            cliente.close()
    if tipo == "application/pdf" or contenido.startswith(b"%PDF-"):
        return _texto_de_pdf(contenido)
    if tipo.endswith("wordprocessingml.document"):
        return _texto_de_docx(contenido)
    texto = _decodificar(contenido)
    if "html" in tipo or re.match(r"\s*<(!doctype|html)", texto, re.IGNORECASE):
        return texto_de_html(texto)
    if tipo.startswith("text/") or tipo in ("", "application/json", "application/xml"):
        return texto
    raise ErrorLectura(f"Tipo de contenido no admitido: {tipo}.")
